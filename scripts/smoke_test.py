from __future__ import annotations

import io
import json
import subprocess
import sys
import tempfile
from dataclasses import replace
from datetime import date, timedelta
from pathlib import Path

from docx import Document

sys.path.insert(0, str(Path(__file__).resolve().parent.parent))

from local_learning_coach.config import Settings
from local_learning_coach.learning import LearningCoachService
from local_learning_coach.interview.runner import CodeRunResult
from local_learning_coach.search import OpenSearchBackend


class MockWebFetcher:
    def fetch(self, url: str, *, check_robots: bool = True):
        import hashlib

        content = (
            b"<!doctype html><html><head><title>Arduino Python Notes</title></head><body>"
            b"<h1>Arduino Python Notes</h1><h2>Variables</h2>"
            b"<p>Python variables and Arduino sensor examples support a practical project.</p>"
            b"</body></html>"
        )
        return {
            "url": url, "content": content, "content_type": "text/html",
            "title": "Arduino Python Notes", "suffix": ".html",
            "sha256": hashlib.sha256(content).hexdigest(),
            "accessed_at": "2026-09-03T00:00:00+00:00",
        }


class MockCodeRunner:
    name = "mock-safe-runner"

    def run(self, code: str, trusted_tests: str, *, approved: bool = False) -> CodeRunResult:
        if not approved:
            raise RuntimeError("Smoke runner açık onay bekler.")
        return CodeRunResult(
            status="passed", passed_tests=2, total_tests=2, elapsed_ms=5, output="",
            code_sha256="0" * 64, runner=self.name,
            isolation_notice="Smoke testinde kullanıcı kodu çalıştırılmadı; güvenli adapter sonucu kullanıldı.",
        )


def command(*args: str) -> str:
    result = subprocess.run(
        [sys.executable, "cli.py", *args],
        capture_output=True,
        text=True,
        encoding="utf-8",
        errors="replace",
        check=False,
    )
    if result.returncode != 0:
        raise RuntimeError(f"Komut başarısız: {' '.join(args)}\n{result.stdout}\n{result.stderr}")
    return result.stdout


def main() -> int:
    if hasattr(sys.stdout, "reconfigure"):
        sys.stdout.reconfigure(encoding="utf-8")
    command("--help")
    command("index", "status")
    command("routes")
    with tempfile.TemporaryDirectory(prefix="llc-smoke-") as temp:
        db_path = Path(temp) / "smoke.db"
        service = LearningCoachService(replace(Settings.load(), db_path=db_path))
        service.setup(build_index=False)
        today = date.today()
        profile_id = service.create_profile(
            {
                "name": "Smoke Test", "goal": "Pandas öğrenmek", "preferred_route": "data-analysis",
                "weekly_days": 5, "daily_minutes": 60,
                "preferred_days": ["Pazartesi", "Salı", "Çarşamba", "Perşembe", "Cuma"],
                "start_date": today.isoformat(), "target_end_date": (today + timedelta(days=120)).isoformat(),
                "theory_practice_preference": "Dengeli", "completed_topics": [], "difficult_topics": [],
                "review_preference": "Düşük sınav puanında", "declared_level": "Başlangıç",
            },
            {"python": "Başlangıç", "pandas": "Yok", "jupyter": "Yok"},
        )
        questions = service.assessment.questions("data-analysis")
        service.assessment.grade(profile_id, "data-analysis", {q.id: q.correct_answer for q in questions})
        plan = service.generate_plan(profile_id)
        tasks = service.database.tasks(profile_id)
        first = tasks[0]
        db_args = ("--db", str(db_path))
        command(*db_args, "today")
        command(*db_args, "task", "start", first["id"])
        command(*db_args, "task", "complete", first["id"], "--minutes", "45", "--difficulty", "3", "--quiz-score", "90")
        second = service.database.tasks(profile_id)[1]
        service.start_task(second["id"])
        completed_second = service.complete_task(second["id"], 50, 4, quiz_score=40)
        command(*db_args, "progress")
        command(*db_args, "weekly-report")
        results = service.retrieve("Titanic planında 12. gün ne yapılıyor?", route="data-analysis", top_k=3)
        if not results or results[0].chunk.get("day") != "12":
            raise RuntimeError("Retrieval smoke testi Day 12 bölümünü bulamadı.")

        service.interview.save_career_profile(profile_id, {
            "department": "Yazılım Mühendisliği", "education_type": "Lisans", "education_status": "Öğrenci",
            "experience": "0-1 yıl", "target_track": "backend-python",
            "target_role": "Junior Backend Engineer", "target_level": "junior", "preferred_language": "Python",
            "other_technologies": ["FastAPI", "PostgreSQL"], "company_type": "Ürün şirketi",
            "interview_language": "Türkçe", "target_interview_date": (today + timedelta(days=45)).isoformat(),
            "weekly_days": 5, "daily_minutes": 60,
            "preferred_days": ["Pazartesi", "Çarşamba", "Cuma"],
            "declared_strengths": ["Python"], "declared_weaknesses": ["Sistem tasarımı"],
            "job_url": "", "job_text": "",
        })
        service.interview.runner = MockCodeRunner()
        interview_session = service.interview.start_assessment(profile_id, seed=20260903)
        while True:
            interview_question = service.interview.current_assessment_question(interview_session["id"])
            if not interview_question:
                break
            if interview_question["question_type"] == "technical":
                answer = interview_question["correct_answer"]
                runner_approved = False
            elif interview_question["question_type"] == "coding":
                answer = "def add(a, b):\n    return a + b"
                runner_approved = True
            else:
                answer = (
                    "set hash O(n) karmaşıklık stack yığın push pop; mutable varsayılan None; JOIN GROUP BY SUM; "
                    "API veri cache hata güvenlik; Situation Task Action Result, somut proje örneğinde ben yaptım "
                    "ve sonuç %20 iyileşti."
                )
                runner_approved = False
            service.interview.submit_assessment_response(
                interview_session["id"], interview_question["id"], answer, runner_approved=runner_approved,
            )
        interview_gap = service.interview.latest_gap_report(profile_id)
        interview_plan = service.interview.create_plan(profile_id, interview_gap["id"])
        cached_interview_plan = service.interview.create_plan(profile_id, interview_gap["id"])
        confirmed_interview_plan = service.interview.confirm_plan(interview_plan["id"], profile_id)
        interview_task = next(item for item in confirmed_interview_plan["tasks"] if item["included"])
        service.interview.update_plan_task_status(interview_task["id"], "in_progress")
        service.interview.update_plan_task_status(interview_task["id"], "completed")
        mock = service.interview.start_mock_interview(profile_id, "Tam karma mülakat")
        behavioral_index = 0
        while True:
            mock_question = service.interview.mock_interview(mock["id"])["current_question"]
            if not mock_question:
                break
            if mock_question["question_type"] == "technical":
                mock_answer, mock_approved = mock_question["correct_answer"], False
            elif mock_question["question_type"] == "coding":
                mock_answer, mock_approved = "def add(a, b):\n    return a + b", True
            elif mock_question["question_type"] == "behavioral" and behavioral_index == 0:
                behavioral_index += 1
                mock_answer, mock_approved = "Kısa yanıt", False
            else:
                mock_answer, mock_approved = (
                    "set hash O(n) karmaşıklık; JOIN GROUP BY SUM; gereksinim trafik API endpoint veri tablo akış "
                    "hata timeout cache TTL ölçek load balancer güven auth rate limit log metrik trace çünkü trade-off; "
                    "Situation Task Action Result somut proje örneğinde ben uyguladım ve sonuç %25 iyileşti.", False
                )
            service.interview.submit_mock_response(
                mock["id"], mock_answer, runner_approved=mock_approved,
            )
        completed_mock = service.interview.mock_action(mock["id"], "finish")
        interview_dashboard = service.interview.dashboard(profile_id)

        custom_data = Path(temp) / "custom-data"
        custom_settings = replace(
            Settings.load(),
            data_dir=custom_data,
            db_path=custom_data / "custom.db",
            index_dir=custom_data / "index",
            embedding_backend="hashing",
            embedding_model="hashing-smoke-v1",
        )
        custom_service = LearningCoachService(custom_settings)
        custom_service.database.migrate()
        custom_profile_id = custom_service.create_profile(
            {
                "name": "Özel Rota Smoke", "goal": "Arduino öğrenmek", "preferred_route": "unassigned",
                "goal_description": "Kendi belgeleri ve onaylı web kaynağıyla öğrenmek",
                "source_method": "Belgeler ve internet kaynaklarını birlikte kullan",
                "weekly_days": 3, "daily_minutes": 45,
                "preferred_days": ["Pazartesi", "Çarşamba", "Cuma"],
                "start_date": today.isoformat(), "target_end_date": (today + timedelta(days=60)).isoformat(),
                "theory_practice_preference": "Dengeli", "completed_topics": [], "difficult_topics": [],
                "review_preference": "Haftalık", "declared_level": "Başlangıç",
            },
            {"python": "Başlangıç"},
        )
        document = Document()
        document.add_heading("Arduino Rehberi", level=1)
        document.add_heading("Dijital Çıkışlar", level=2)
        document.add_paragraph("Arduino dijital pinleri LED ve benzeri çıkışları denetlemek için kullanılabilir.")
        document.add_heading("Analog Sensörler", level=2)
        document.add_paragraph("Analog sensör değerleri okunur ve seri monitörde gözlemlenir.")
        buffer = io.BytesIO()
        document.save(buffer)
        uploaded = custom_service.upload_document("Arduino Smoke.docx", buffer.getvalue())
        document2 = Document()
        document2.add_heading("Arduino Uygulamaları", level=1)
        document2.add_heading("Sensör Projesi", level=2)
        document2.add_paragraph("Arduino analog sensor measurements are used in a practical example project.")
        buffer2 = io.BytesIO()
        document2.save(buffer2)
        uploaded2 = custom_service.upload_document("Arduino Project Smoke.docx", buffer2.getvalue())
        discovered = custom_service.add_web_url("https://93.184.216.34/arduino", title="Arduino Python Notes")
        approved_web = custom_service.web_sources.approve(discovered["id"], fetcher=MockWebFetcher())
        job_requirements = custom_service.interview.extract_job_requirements([approved_web["document_id"]])
        configuration = {
            "title": "Arduino Smoke Rotası", "goal": "Arduino analog sensor Python variables project",
            "declared_level": "Başlangıç", "weekly_days": 3, "daily_minutes": 45,
            "preferred_days": ["Pazartesi", "Çarşamba", "Cuma"],
            "start_date": today.isoformat(), "target_end_date": (today + timedelta(days=60)).isoformat(),
            "theory_practice_preference": "Dengeli",
        }
        source_ids = [uploaded["id"], uploaded2["id"], approved_web["document_id"]]
        sufficiency = custom_service.analyze_sources(custom_profile_id, source_ids, configuration)
        if not sufficiency["can_generate"]:
            raise RuntimeError("Kaynak yeterlilik smoke testi rota üretimine izin vermedi.")
        custom_service.approve_sufficiency(sufficiency["id"])
        draft = custom_service.create_multi_source_route_draft(
            source_ids, custom_profile_id, configuration, sufficiency["id"]
        )
        custom_quiz = custom_service.create_custom_quiz(draft["route"]["id"], source_ids)
        custom_plan = custom_service.install_custom_route(draft["route"]["id"], custom_profile_id)
        custom_results = custom_service.retrieve("analog sensör", document_id=uploaded["id"])
        custom_first = custom_service.database.tasks(custom_profile_id)[0]
        custom_service.focus_action("start", custom_profile_id, custom_first["id"])
        custom_service.focus_action("pause", custom_profile_id)
        custom_service.focus_action("resume", custom_profile_id)
        focus = custom_service.focus_action("finish", custom_profile_id)
        custom_service.start_task(custom_first["id"])
        completed_custom = custom_service.complete_task(custom_first["id"], 60, 4, quiz_score=40)
        analytics = custom_service.analytics(custom_profile_id)
        opensearch_mock_contract = OpenSearchBackend(custom_settings).name == "opensearch"
        if not custom_plan["task_count"] or not custom_results or not custom_quiz:
            raise RuntimeError("Özel belge ve rota smoke testi başarısız oldu.")
        payload = {
            "status": "passed",
            "cli": ["--help", "index status", "routes", "today", "task start", "task complete", "progress", "weekly-report"],
            "profile_id": profile_id,
            "plan_tasks": plan["task_count"],
            "review_task_created": bool(completed_second.get("review_task")),
            "retrieval_top": results[0].chunk["section_path"],
            "progress": service.progress(profile_id),
            "custom_document": uploaded["original_filename"],
            "multi_source_count": len(source_ids),
            "web_snapshot": approved_web["snapshot_path"],
            "sufficiency_score": sufficiency["coverage_score"],
            "custom_plan_tasks": custom_plan["task_count"],
            "custom_quiz_questions": len(custom_quiz),
            "focus_elapsed_seconds": focus["elapsed_seconds"],
            "adaptive_review_created": bool(completed_custom.get("review_task")),
            "analytics_task_rows": len(analytics["tasks"]),
            "search_backend": custom_service.search_backend_status()["backend"],
            "mock_opensearch_contract": opensearch_mock_contract,
            "custom_retrieval_top": custom_results[0].chunk["section_path"],
            "interview_questions": len(interview_session["question_ids"]),
            "interview_gap_confidence": interview_gap["confidence_score"],
            "interview_plan_tasks": len(confirmed_interview_plan["tasks"]),
            "interview_task_completed": True,
            "mock_interview_status": completed_mock["status"],
            "mock_question_types": sorted({
                item["question_type"] for item in service.interview._questions(mock["track_id"], mock["question_ids"])
            }),
            "adaptive_decisions": len(interview_dashboard["adaptive_decisions"]),
            "dashboard_target": interview_dashboard["target_role"],
            "plan_cache_hit": cached_interview_plan["cache_hit"],
            "job_skills": [item["skill"] for item in job_requirements["skills"]],
        }
    print(json.dumps(payload, ensure_ascii=False, indent=2))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
