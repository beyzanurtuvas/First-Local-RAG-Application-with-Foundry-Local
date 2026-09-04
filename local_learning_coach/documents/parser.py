from __future__ import annotations

import hashlib
import json
import re
import unicodedata
from dataclasses import asdict, dataclass
from html.parser import HTMLParser
from pathlib import Path
from typing import Iterable, Iterator

from docx import Document
from docx.document import Document as DocumentType
from docx.oxml.ns import qn
from docx.table import Table
from docx.text.paragraph import Paragraph
from pypdf import PdfReader

from local_learning_coach.config import route_for_filename
from local_learning_coach.documents.ocr import OCRAdapter


SPACE_RE = re.compile(r"\s+")
PHASE_RE = re.compile(r"(?i)\bPhase\s+(\d+)(?:\s*[:(–—-]|\b)")
WEEK_RE = re.compile(r"(?i)^Week\s+(\d+)(?:\s*[:(–—-]|\b)")
DAY_RE = re.compile(r"(?i)^Days?\s+(\d+)(?:\s*[–—-]\s*(\d+))?(?:\s*:|\b)")


def clean_text(value: str) -> str:
    return SPACE_RE.sub(" ", unicodedata.normalize("NFC", value)).strip()


def stable_hash(*parts: str, length: int = 24) -> str:
    normalized = "\x1f".join(clean_text(part).casefold() for part in parts)
    return hashlib.sha256(normalized.encode("utf-8")).hexdigest()[:length]


@dataclass(frozen=True)
class Block:
    kind: str
    text: str
    style: str
    bold_heading: bool
    list_kind: str | None
    urls: tuple[str, ...]
    page: int | None = None
    extraction_method: str = "native"
    ocr_confidence: float | None = None


@dataclass(frozen=True)
class Chunk:
    id: str
    document_id: str
    source: str
    doc_title: str
    route: str
    section_path: str
    heading: str
    phase: str | None
    week: str | None
    day: str | None
    page: int | None
    body: str
    text: str
    urls: tuple[str, ...]
    n_chars: int
    order: int
    content_hash: str
    extraction_method: str = "native"
    ocr_confidence: float | None = None

    def to_dict(self) -> dict:
        result = asdict(self)
        result["urls"] = list(self.urls)
        return result


@dataclass(frozen=True)
class ParsedDocument:
    id: str
    filename: str
    path: str
    title: str
    route: str
    route_title: str
    sha256: str
    chunks: tuple[Chunk, ...]
    display_name: str | None = None

    def metadata(self) -> dict:
        return {
            "id": self.id,
            "filename": self.filename,
            "path": self.path,
            "title": self.title,
            "route": self.route,
            "route_title": self.route_title,
            "sha256": self.sha256,
            "original_filename": self.display_name or self.filename,
        }


@dataclass(frozen=True)
class SourceTask:
    source_document: str
    source_section: str
    source_chunk_id: str
    route: str
    phase: str | None
    week: str | None
    day: str | None
    title: str
    description: str
    task_type: str
    estimated_minutes: int
    estimated_is_system: bool
    position: int
    is_critical: bool
    provenance: str = "belgeden çıkarıldı"
    source_page: int | None = None
    planning_key: str | None = None
    module: str | None = None
    confidence_explanation: str = ""
    generation_type: str = ""
    prerequisite_planning_keys: tuple[str, ...] = ()


def _iter_blocks(document: DocumentType) -> Iterator[tuple[str, Paragraph | Table]]:
    paragraphs = {p._p: p for p in document.paragraphs}
    tables = {t._tbl: t for t in document.tables}
    for child in document.element.body.iterchildren():
        if child.tag == qn("w:p"):
            yield "paragraph", paragraphs[child]
        elif child.tag == qn("w:tbl"):
            yield "table", tables[child]


def _paragraph_urls(paragraph: Paragraph) -> tuple[str, ...]:
    urls: list[str] = []
    for hyperlink in paragraph._p.xpath(".//w:hyperlink"):
        rel_id = hyperlink.get(qn("r:id"))
        if rel_id and rel_id in paragraph.part.rels:
            urls.append(str(paragraph.part.rels[rel_id].target_ref))
    urls.extend(match.rstrip(".,;]") for match in re.findall(r"https?://[^\s)>]+", paragraph.text))
    return tuple(dict.fromkeys(urls))


def _list_kind(paragraph: Paragraph) -> str | None:
    p_pr = paragraph._p.pPr
    if p_pr is None or p_pr.numPr is None:
        return None
    ilvl = p_pr.numPr.ilvl
    num_id = p_pr.numPr.numId
    return f"numbered-level-{ilvl.val if ilvl is not None else 0}-id-{num_id.val if num_id is not None else 0}"


def _is_bold_heading(paragraph: Paragraph, text: str) -> bool:
    meaningful = [run for run in paragraph.runs if clean_text(run.text)]
    if not meaningful:
        return False
    if all(run.bold for run in meaningful):
        return len(text) <= 180 or text.endswith(":")
    leading = clean_text(meaningful[0].text)
    return bool(meaningful[0].bold and leading.endswith(":") and len(leading) <= 120)


def _table_block(table: Table) -> Block | None:
    rows = [[clean_text(cell.text) for cell in row.cells] for row in table.rows]
    rows = [row for row in rows if any(row)]
    if not rows:
        return None
    width = max(len(row) for row in rows)
    headers = [(rows[0][i] if i < len(rows[0]) and rows[0][i] else f"Sütun {i + 1}") for i in range(width)]
    lines: list[str] = []
    for row_index, row in enumerate(rows[1:] or rows):
        values = rows[row_index + 1] if len(rows) > 1 else row
        cells = [f"{headers[i]}: {values[i]}" for i in range(min(width, len(values))) if values[i]]
        if cells:
            lines.append(" | ".join(cells))
    urls = tuple(
        dict.fromkeys(
            url
            for row in table.rows
            for cell in row.cells
            for paragraph in cell.paragraphs
            for url in _paragraph_urls(paragraph)
        )
    )
    return Block("table", "\n".join(lines), "Table", False, None, urls)


def _extract_heading_and_body(text: str, marker: re.Match[str]) -> tuple[str, str]:
    after = text[marker.end():].lstrip(" :–—-")
    marker_text = clean_text(text[: marker.end()].rstrip(" :–—-"))
    if not after:
        return clean_text(text), ""
    if ":" in after:
        label, body = after.split(":", 1)
        if 1 <= len(label) <= 110:
            return f"{marker_text}: {clean_text(label)}", clean_text(body)
    if len(after) <= 120:
        return clean_text(text), ""
    return marker_text, after


class DocumentParser:
    def __init__(
        self,
        max_chars: int = 2400,
        overlap_paragraphs: int = 1,
        ocr_adapter: OCRAdapter | None = None,
    ):
        if max_chars < 500:
            raise ValueError("chunk max_chars en az 500 olmalıdır")
        if overlap_paragraphs < 0 or overlap_paragraphs > 3:
            raise ValueError("overlap_paragraphs 0 ile 3 arasında olmalıdır")
        self.max_chars = max_chars
        self.overlap_paragraphs = overlap_paragraphs
        self.ocr_adapter = ocr_adapter

    def read_blocks(self, path: Path) -> list[Block]:
        if path.suffix.casefold() == ".pdf":
            return self._read_pdf_blocks(path)
        if path.suffix.casefold() in {".html", ".htm", ".txt"}:
            return self._read_web_blocks(path)
        if path.suffix.casefold() != ".docx":
            raise ValueError(f"Desteklenmeyen belge türü: {path.suffix}")
        document = Document(path)
        blocks: list[Block] = []
        for kind, item in _iter_blocks(document):
            if kind == "table":
                block = _table_block(item)  # type: ignore[arg-type]
                if block:
                    blocks.append(block)
                continue
            paragraph: Paragraph = item  # type: ignore[assignment]
            text = clean_text(paragraph.text)
            if not text:
                continue
            style = paragraph.style.name if paragraph.style else ""
            blocks.append(
                Block(
                    "paragraph",
                    text,
                    style,
                    _is_bold_heading(paragraph, text),
                    _list_kind(paragraph),
                    _paragraph_urls(paragraph),
                )
            )
        return blocks

    def _read_pdf_blocks(self, path: Path) -> list[Block]:
        reader = PdfReader(path)
        if reader.is_encrypted and reader.decrypt("") == 0:
            raise ValueError("Parolalı PDF belgeleri desteklenmiyor.")
        if not reader.pages:
            raise ValueError("PDF belgesinde sayfa bulunamadı.")
        native_pages: list[str] = []
        for page in reader.pages:
            try:
                native_pages.append(page.extract_text(extraction_mode="layout") or "")
            except (TypeError, ValueError):
                native_pages.append(page.extract_text() or "")
        extraction_method = "native"
        confidences: dict[int, float | None] = {}
        if sum(len(clean_text(value)) for value in native_pages) < 40 and self.ocr_adapter is not None:
            ocr_pages = self.ocr_adapter.extract_pdf(path)
            native_pages = [""] * len(reader.pages)
            for item in ocr_pages:
                if 1 <= item.page <= len(native_pages):
                    native_pages[item.page - 1] = item.text
                    confidences[item.page] = item.confidence
            extraction_method = "ocr"
        blocks: list[Block] = []
        first_text = True
        for page_number, text in enumerate(native_pages, start=1):
            for raw_line in text.splitlines():
                table_like = bool(re.search(r"\S\s{2,}\S", raw_line.strip()))
                preserved_line = re.sub(r"\s{2,}", " | ", raw_line.strip()) if table_like else raw_line
                line = clean_text(preserved_line)
                if not line:
                    continue
                is_list = bool(re.match(r"^(?:[•●▪*-]|\d+[.)])\s+", line))
                heading_like = (
                    first_text
                    or (
                        len(line) <= 120
                        and len(line.split()) <= 12
                        and not line.endswith((".", "?", "!", ";"))
                        and bool(re.match(r"^(?:\d+(?:\.\d+)*[.)]?\s+)?[A-ZÇĞİÖŞÜ]", line))
                    )
                    or (len(line) <= 100 and line.isupper())
                )
                style = "Heading 1" if first_text else (
                    "Heading 2" if heading_like and not is_list and not table_like else ("PDF Table" if table_like else "PDF Text")
                )
                blocks.append(
                    Block(
                        "table" if table_like else "paragraph",
                        line,
                        style,
                        heading_like and not table_like,
                        "pdf-list" if is_list else None,
                        (),
                        page_number,
                        extraction_method,
                        confidences.get(page_number),
                    )
                )
                first_text = False
        return blocks

    @staticmethod
    def _read_web_blocks(path: Path) -> list[Block]:
        text = path.read_text(encoding="utf-8", errors="replace")
        if path.suffix.casefold() == ".txt":
            return [
                Block("paragraph", clean_text(line), "Web Text", False, None, (), None, "web")
                for line in text.splitlines()
                if clean_text(line)
            ]

        class Reader(HTMLParser):
            def __init__(self) -> None:
                super().__init__(convert_charrefs=True)
                self.tag = ""
                self.data: list[str] = []
                self.blocks: list[Block] = []
                self.hrefs: list[str] = []

            def handle_starttag(self, tag: str, attrs: list[tuple[str, str | None]]) -> None:
                if tag in {"h1", "h2", "h3", "h4", "p", "li", "pre", "code", "tr", "blockquote"}:
                    self.flush()
                    self.tag = tag
                if tag == "a":
                    href = dict(attrs).get("href")
                    if href:
                        self.hrefs.append(href)

            def handle_endtag(self, tag: str) -> None:
                if tag == self.tag:
                    self.flush()

            def handle_data(self, data: str) -> None:
                if self.tag:
                    self.data.append(data)

            def flush(self) -> None:
                value = clean_text(" ".join(self.data))
                if value and self.tag:
                    heading = self.tag.startswith("h")
                    kind = "table" if self.tag == "tr" else "paragraph"
                    style = f"Heading {min(int(self.tag[1]), 3)}" if heading else (
                        "Web Code" if self.tag in {"pre", "code"} else "Web Text"
                    )
                    self.blocks.append(
                        Block(kind, value, style, heading, "web-list" if self.tag == "li" else None,
                              tuple(dict.fromkeys(self.hrefs + re.findall(r"https?://[^\s\])}]+", value))), None, "web")
                    )
                self.tag = ""
                self.data = []
                self.hrefs = []

        reader = Reader()
        reader.feed(text)
        reader.flush()
        return reader.blocks

    def parse(
        self,
        path: Path,
        *,
        route: str | None = None,
        route_title: str | None = None,
        document_id: str | None = None,
        source_name: str | None = None,
    ) -> ParsedDocument:
        if route is None or route_title is None:
            route, route_title = route_for_filename(path.name)
        raw = path.read_bytes()
        sha256 = hashlib.sha256(raw).hexdigest()
        resolved_document_id = document_id or stable_hash(path.name, length=20)
        display_name = source_name or path.name
        blocks = self.read_blocks(path)
        if not blocks:
            raise ValueError(f"Belge boş veya okunamıyor: {path}")
        title_block = next((b for b in blocks if b.style.lower().startswith("heading 1")), blocks[0])
        doc_title = title_block.text

        sections: list[dict] = []
        current = {
            "heading": doc_title,
            "phase": None,
            "week": None,
            "day": None,
            "page": title_block.page,
            "paragraphs": [],
            "urls": [],
            "extraction_method": title_block.extraction_method,
            "ocr_confidence": title_block.ocr_confidence,
            "path": [doc_title] + ([f"Sayfa {title_block.page}"] if title_block.page else []),
        }
        phase: str | None = None
        week: str | None = None
        day: str | None = None

        def flush() -> None:
            nonlocal current
            if current["paragraphs"]:
                sections.append(current)

        for block in blocks:
            if block is title_block:
                continue
            text = block.text
            phase_match = PHASE_RE.search(text)
            week_match = WEEK_RE.search(text)
            day_match = DAY_RE.search(text)
            is_style_heading = block.style.lower().startswith("heading")
            structural = bool(phase_match or week_match or day_match)
            bold_section = block.bold_heading and text.casefold() not in {"activities:", "activity:", "etkinlikler:"}
            page_changed = block.page is not None and block.page != current.get("page")
            split_section = page_changed or structural or (is_style_heading and block.style.lower() != "heading 1") or bold_section

            heading = text
            body = ""
            if day_match:
                heading, body = _extract_heading_and_body(text, day_match)
            elif week_match and not phase_match:
                heading, body = _extract_heading_and_body(text, week_match)
            elif bold_section and not structural and ":" in text:
                label, remainder = text.split(":", 1)
                if len(label) <= 120:
                    heading = clean_text(label) + ":"
                    body = clean_text(remainder)
            elif page_changed and not is_style_heading:
                heading = f"Sayfa {block.page}"
                body = text

            if split_section:
                flush()
                if phase_match:
                    phase = phase_match.group(1)
                    week = None
                    day = None
                if week_match and not phase_match:
                    week = week_match.group(1)
                    day = None
                if day_match:
                    start_day = day_match.group(1)
                    end_day = day_match.group(2)
                    day = f"{start_day}-{end_day}" if end_day else start_day
                path_parts = [doc_title]
                if phase:
                    path_parts.append(f"Phase {phase}")
                if week:
                    path_parts.append(f"Week {week}")
                if day:
                    path_parts.append(f"Day {day}")
                if block.page:
                    path_parts.append(f"Sayfa {block.page}")
                if heading not in path_parts:
                    path_parts.append(heading)
                current = {
                    "heading": heading,
                    "phase": phase,
                    "week": week,
                    "day": day,
                    "page": block.page,
                    "paragraphs": [],
                    "urls": list(block.urls),
                    "extraction_method": block.extraction_method,
                    "ocr_confidence": block.ocr_confidence,
                    "path": path_parts,
                }
                if body:
                    current["paragraphs"].append(body)
                continue

            prefix = ""
            if block.kind == "table":
                prefix = "Tablo:\n"
            elif block.list_kind:
                prefix = "• "
            current["paragraphs"].append(prefix + text)
            current["urls"].extend(block.urls)
        flush()

        chunks: list[Chunk] = []
        order = 0
        for section in sections:
            for part in self._split_section(section["paragraphs"]):
                body = "\n\n".join(part).strip()
                if len(body) < 30 and chunks and section["heading"] == chunks[-1].heading:
                    continue
                section_path = " > ".join(section["path"])
                content_hash = hashlib.sha256(clean_text(body).encode("utf-8")).hexdigest()
                chunk_identity = path.name if document_id is None else resolved_document_id
                chunk_id = stable_hash(chunk_identity, section_path, body, length=24)
                compact_section = " > ".join(section["path"][1:]) or section["heading"]
                embedding_text = (
                    f"Document: {display_name}\nRoute: {route_title}\nSection: {compact_section}\n{body}"
                )
                chunks.append(
                    Chunk(
                        id=chunk_id,
                        document_id=resolved_document_id,
                        source=display_name,
                        doc_title=doc_title,
                        route=route,
                        section_path=section_path,
                        heading=section["heading"],
                        phase=section["phase"],
                        week=section["week"],
                        day=section["day"],
                        page=section.get("page"),
                        body=body,
                        text=embedding_text,
                        urls=tuple(dict.fromkeys(section["urls"])),
                        n_chars=len(embedding_text),
                        order=order,
                        content_hash=content_hash,
                        extraction_method=section.get("extraction_method", "native"),
                        ocr_confidence=section.get("ocr_confidence"),
                    )
                )
                order += 1
        return ParsedDocument(
            id=resolved_document_id,
            filename=path.name,
            path=str(path),
            title=doc_title,
            route=route,
            route_title=route_title,
            sha256=sha256,
            chunks=tuple(chunks),
            display_name=display_name,
        )

    def parse_all(self, paths: Iterable[Path]) -> list[ParsedDocument]:
        return [self.parse(path) for path in paths]

    def _split_section(self, paragraphs: list[str]) -> list[list[str]]:
        if not paragraphs:
            return []
        parts: list[list[str]] = []
        current: list[str] = []
        current_len = 0
        for paragraph in paragraphs:
            if len(paragraph) > self.max_chars:
                sentences = re.split(r"(?<=[.!?])\s+", paragraph)
            else:
                sentences = [paragraph]
            for unit in sentences:
                added = len(unit) + (2 if current else 0)
                if current and current_len + added > self.max_chars:
                    parts.append(current)
                    current = current[-self.overlap_paragraphs :] if self.overlap_paragraphs else []
                    current_len = sum(len(item) + 2 for item in current)
                current.append(unit)
                current_len += added
        if current:
            parts.append(current)
        return parts


def task_type_for(text: str) -> str:
    lowered = text.casefold()
    if any(word in lowered for word in ("quiz", "assessment", "checkpoint")):
        return "değerlendirme"
    if any(word in lowered for word in ("document", "readme", "report", "presentation")):
        return "dokümantasyon"
    if any(word in lowered for word in ("test", "validate", "evaluation", "evaluate", "metric")):
        return "test ve değerlendirme"
    if any(word in lowered for word in ("implement", "project", "code", "notebook", "exercise", "practice")):
        return "uygulama"
    return "teori"


def extract_source_tasks(document: ParsedDocument) -> list[SourceTask]:
    has_days = any(chunk.day for chunk in document.chunks)
    grouped: dict[tuple[str | None, str | None, str | None], list[Chunk]] = {}
    for chunk in document.chunks:
        if has_days:
            if not chunk.day:
                continue
            key = (chunk.phase, chunk.week, chunk.day)
        else:
            if not chunk.week:
                continue
            key = (chunk.phase, chunk.week, None)
        grouped.setdefault(key, []).append(chunk)
    tasks: list[SourceTask] = []
    for position, (_, chunks) in enumerate(grouped.items(), start=1):
        first = chunks[0]
        bodies = list(dict.fromkeys(chunk.body for chunk in chunks if chunk.body))
        description = "\n\n".join(bodies)
        # Belgelerde süre bulunmadığında şeffaf, deterministik bir tahmin kullanılır.
        estimated = min(300, max(45, ((len(description) // 900) + 1) * 45))
        title = first.heading
        tasks.append(
            SourceTask(
                source_document=first.source,
                source_section=first.section_path,
                source_chunk_id=first.id,
                route=document.route,
                phase=first.phase,
                week=first.week,
                day=first.day,
                title=title,
                description=description,
                task_type=task_type_for(title + " " + description),
                estimated_minutes=estimated,
                estimated_is_system=True,
                position=position,
                is_critical=position == 1 or "milestone" in description.casefold(),
                provenance="belgeden çıkarıldı",
                source_page=first.page,
            )
        )
    return tasks


def write_chunks_jsonl(documents: Iterable[ParsedDocument], path: Path) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    with path.open("w", encoding="utf-8", newline="\n") as handle:
        for document in documents:
            for chunk in document.chunks:
                handle.write(json.dumps(chunk.to_dict(), ensure_ascii=False) + "\n")
