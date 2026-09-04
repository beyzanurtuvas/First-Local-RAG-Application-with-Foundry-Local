from __future__ import annotations

import io
from datetime import date, timedelta
from pathlib import Path

import pytest
from docx import Document

from local_learning_coach.database import DatabaseError
from local_learning_coach.learning import LearningCoachService, UploadError


def docx_bytes(*, structured: bool = False) -> bytes:
    document = Document()
    document.add_heading("Arduino Başlangıç Rehberi", level=1)
    if structured:
        document.add_heading("Week 1: Elektronik Temelleri", level=2)
        document.add_paragraph("Arduino kartını, dijital pinleri ve güvenli devre kurulumunu öğrenin.")
        document.add_heading("Week 2: Sensör Uygulaması", level=2)
        document.add_paragraph("Bir sensörden ölçüm alın ve seri monitörde gözlemleyin.")
    else:
        document.add_heading("Elektronik ve Kart Yapısı", level=2)
        document.add_paragraph(
            "Arduino kartı mikrodenetleyici, dijital pinler ve analog girişler içerir. "
            "Devre kurulurken uygun direnç ve ortak toprak bağlantısı kullanılmalıdır."
        )
        document.add_heading("Dijital Giriş ve Çıkış", level=2)
        document.add_paragraph(
            "Dijital pinler giriş veya çıkış olarak ayarlanabilir. LED kontrolü temel bir çıkış örneğidir."
        )
        document.add_heading("Sensörlerden Veri Okuma", level=2)
        document.add_paragraph(
            "Analog sensör değerleri okunabilir ve seri monitör üzerinden incelenebilir. Ölçümler tekrar edilmelidir."
        )
        table = document.add_table(rows=2, cols=2)
        table.rows[0].cells[0].text = "Bileşen"
        table.rows[0].cells[1].text = "Amaç"
        table.rows[1].cells[0].text = "LED"
        table.rows[1].cells[1].text = "Dijital çıkışı gözlemlemek"
    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def pdf_bytes() -> bytes:
    stream = b"BT /F1 16 Tf 72 740 Td (Robotik Temelleri) Tj /F1 12 Tf 0 -30 Td (Motor Kontrolu) Tj 0 -20 Td (DC motorlar surucu devre ile kontrol edilir.) Tj ET"
    objects = [
        b"<< /Type /Catalog /Pages 2 0 R >>",
        b"<< /Type /Pages /Kids [3 0 R] /Count 1 >>",
        b"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] /Resources << /Font << /F1 5 0 R >> >> /Contents 4 0 R >>",
        b"<< /Length %d >>\nstream\n%s\nendstream" % (len(stream), stream),
        b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>",
    ]
    payload = bytearray(b"%PDF-1.4\n")
    offsets = [0]
    for number, obj in enumerate(objects, start=1):
        offsets.append(len(payload))
        payload.extend(f"{number} 0 obj\n".encode())
        payload.extend(obj + b"\nendobj\n")
    xref = len(payload)
    payload.extend(f"xref\n0 {len(objects) + 1}\n".encode())
    payload.extend(b"0000000000 65535 f \n")
    for offset in offsets[1:]:
        payload.extend(f"{offset:010d} 00000 n \n".encode())
    payload.extend(
        f"trailer\n<< /Size {len(objects) + 1} /Root 1 0 R >>\nstartxref\n{xref}\n%%EOF\n".encode()
    )
    return bytes(payload)


def route_configuration() -> dict:
    return {
        "title": "Arduino ile Uygulamalı Öğrenme",
        "goal": "Arduino ile güvenli ve küçük projeler geliştirmek",
        "declared_level": "Başlangıç",
        "weekly_days": 5,
        "daily_minutes": 60,
        "preferred_days": ["Pazartesi", "Salı", "Çarşamba", "Perşembe", "Cuma"],
        "start_date": date.today().isoformat(),
        "target_end_date": (date.today() + timedelta(days=60)).isoformat(),
        "theory_practice_preference": "Dengeli",
    }


def service_for(test_settings) -> LearningCoachService:
    service = LearningCoachService(test_settings)
    service.database.migrate()
    return service


def test_docx_upload_duplicate_and_retrieval(test_settings):
    service = service_for(test_settings)
    first = service.upload_document("Arduino Rehberi.docx", docx_bytes())
    assert first["origin"] == "uploaded" and first["status"] == "active"
    assert first["chunk_count"] >= 3
    assert Path(first["path"]).parent == test_settings.data_dir / "uploads"
    assert Path(first["path"]).name.startswith("uploaded-")

    duplicate = service.upload_document("kopya.docx", docx_bytes())
    assert duplicate["duplicate"] is True and duplicate["id"] == first["id"]
    assert len(service.uploaded_documents()) == 1

    results = service.retrieve("analog sensör değerleri", document_id=first["id"])
    assert results and all(item.chunk["document_id"] == first["id"] for item in results)
    assert "sensör" in results[0].chunk["body"].casefold()


@pytest.mark.parametrize("filename", ["../kaçış.docx", "..\\kaçış.docx", "C:\\kaçış.docx"])
def test_unsafe_filename_is_rejected(test_settings, filename):
    service = service_for(test_settings)
    with pytest.raises(UploadError, match="Dosya adı"):
        service.upload_document(filename, docx_bytes())
    assert not service.uploaded_documents()


def test_extension_and_signature_must_match(test_settings):
    service = service_for(test_settings)
    with pytest.raises(UploadError, match="geçerli bir PDF"):
        service.upload_document("sahte.pdf", b"not-a-pdf")


def test_pdf_upload_preserves_page_metadata(test_settings):
    service = service_for(test_settings)
    uploaded = service.upload_document("Robotik.pdf", pdf_bytes())
    results = service.retrieve("DC motorlar", document_id=uploaded["id"])
    assert results and results[0].chunk["page"] == 1
    with service.database.connection() as conn:
        pages = [row[0] for row in conn.execute("SELECT page FROM source_chunks WHERE document_id=?", (uploaded["id"],))]
    assert pages and set(pages) == {1}


def test_unstructured_document_route_edit_install_and_progress(test_settings, profile_id):
    service = service_for(test_settings)
    uploaded = service.upload_document("Arduino Rehberi.docx", docx_bytes())
    configuration = {
        **route_configuration(),
        "weekly_days": 3,
        "daily_minutes": 45,
        "preferred_days": ["Pazartesi", "Çarşamba", "Cuma"],
    }
    draft = service.create_custom_route_draft(uploaded["id"], profile_id, configuration)
    tasks = draft["tasks"]
    assert tasks
    assert all(item["source_chunk_id"] and item["source_section"] for item in tasks)
    assert {item["provenance"] for item in tasks} <= {
        "belgeden çıkarıldı", "sistem tarafından oluşturulan alıştırma", "belge dışı ön koşul önerisi"
    }
    assert any(item["provenance"] == "sistem tarafından oluşturulan alıştırma" for item in tasks)

    edits = []
    for index, item in enumerate(reversed(tasks), start=1):
        edits.append(
            {
                **item,
                "position": index,
                "included": item["id"] != tasks[0]["id"],
                "title": "Düzenlendi: " + item["title"] if index == 1 else item["title"],
                "estimated_minutes": 30 if index == 1 else item["estimated_minutes"],
            }
        )
    service.update_custom_route_tasks(draft["route"]["id"], edits)
    edited = service.custom_route_tasks(draft["route"]["id"])
    assert edited[0]["title"].startswith("Düzenlendi:")
    assert edited[0]["estimated_minutes"] == 30
    assert sum(item["included"] for item in edited) == len(tasks) - 1

    installed = service.install_custom_route(draft["route"]["id"], profile_id)
    plan_tasks = service.database.tasks(profile_id)
    assert installed["task_count"] == len(plan_tasks) == len(tasks) - 1
    assert all(item["route"] == draft["route"]["id"] for item in plan_tasks)
    assert all(item["source_chunk_id"] and item["provenance"] for item in plan_tasks)
    updated_profile = service.database.get_profile(profile_id)
    assert updated_profile["preferred_route"] == draft["route"]["id"]
    assert updated_profile["weekly_days"] == 3 and updated_profile["daily_minutes"] == 45
    assert updated_profile["preferred_days"] == ["Pazartesi", "Çarşamba", "Cuma"]

    service.start_task(plan_tasks[0]["id"])
    service.complete_task(plan_tasks[0]["id"], 25, 2)
    assert service.progress(profile_id)["completed"] == 1


def test_structured_document_preserves_week_order(test_settings, profile_id):
    service = service_for(test_settings)
    uploaded = service.upload_document("Haftalık Arduino.docx", docx_bytes(structured=True))
    draft = service.create_custom_route_draft(uploaded["id"], profile_id, route_configuration())
    tasks = draft["tasks"]
    assert [item["week"] for item in tasks] == ["1", "2"]
    assert all(item["provenance"] == "belgeden çıkarıldı" for item in tasks)


def test_archive_requires_exact_confirmation_and_keeps_file(test_settings):
    service = service_for(test_settings)
    uploaded = service.upload_document("Arduino Rehberi.docx", docx_bytes())
    with pytest.raises(DatabaseError, match="Onay metni"):
        service.archive_uploaded_document(uploaded["id"], "evet")
    archived = service.archive_uploaded_document(
        uploaded["id"], f"BELGEYİ ARŞİVLE {uploaded['id']}"
    )
    assert archived["status"] == "archived"
    assert Path(archived["path"]).is_file()
    assert not service.retrieve("Arduino", document_id=uploaded["id"])


def test_builtin_documents_cannot_be_archived(test_settings):
    service = service_for(test_settings)
    service.build_index(force=True)
    builtin = service.database.list_source_documents(origin="builtin")[0]
    with pytest.raises(DatabaseError, match="Hazır üç"):
        service.database.archive_uploaded_document(
            builtin["id"], f"BELGEYİ ARŞİVLE {builtin['id']}"
        )
