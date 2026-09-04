from __future__ import annotations

import os
from dataclasses import replace
from datetime import date, timedelta
from pathlib import Path

import pytest
from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

from local_learning_coach.config import DOCUMENT_DEFINITIONS, Settings
from local_learning_coach.database import Database


def add_hyperlink(paragraph, text: str, url: str) -> None:
    relationship = paragraph.part.relate_to(
        url,
        "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink",
        is_external=True,
    )
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("r:id"), relationship)
    run = OxmlElement("w:r")
    text_element = OxmlElement("w:t")
    text_element.text = text
    run.append(text_element)
    hyperlink.append(run)
    paragraph._p.append(hyperlink)


def make_doc(path: Path, route: str) -> None:
    doc = Document()
    if route == "machine-learning":
        doc.add_heading("Synthetic ML Plan", level=1)
        doc.add_heading("Phase 1: Foundations (Weeks 1–2)", level=2)
        doc.add_paragraph("Week 1: ML basics and regression.")
        doc.add_paragraph("Day 1: Setup: Install Python and Jupyter.")
        doc.add_paragraph("Day 2: Metrics: Learn MSE and RMSE.")
        doc.add_heading("Phase 2: PyTorch (Weeks 3–4)", level=2)
        doc.add_paragraph("Week 3: PyTorch fundamentals and RNN concepts.")
        doc.add_paragraph("Day 11–12: PyTorch Quickstart: Tensors and training loop.")
        doc.add_paragraph("Day 16–17: Data preparation: Sliding window for stock time series.")
        table = doc.add_table(rows=2, cols=2)
        table.rows[0].cells[0].text = "Week"
        table.rows[0].cells[1].text = "Outcome"
        table.rows[1].cells[0].text = "3"
        table.rows[1].cells[1].text = "PyTorch begins"
    elif route == "data-analysis":
        doc.add_heading("Synthetic Titanic Plan", level=1)
        doc.add_heading("Phase 1: Python and Pandas", level=2)
        doc.add_paragraph("Day 5: Intro to Pandas: Read CSV into DataFrames.")
        doc.add_paragraph("Day 8: Grouping and Missing Data: GroupBy and cleaning.")
        doc.add_paragraph("Day 9: Data Visualization: Matplotlib and Seaborn.")
        doc.add_heading("Phase 2: Titanic EDA", level=2)
        doc.add_paragraph("Day 12: Travel & Family: Analyze cabins, embarkation, and companions.")
        bold = doc.add_paragraph()
        run = bold.add_run("Deliverable:")
        run.bold = True
        bold.add_run(" Complete the notebook.")
        link = doc.add_paragraph("Resource: ")
        add_hyperlink(link, "example", "https://example.com/titanic")
    else:
        doc.add_heading("Synthetic Quantum Plan", level=1)
        doc.add_heading("Phase 1 (Weeks 1–2): Q# Basics", level=2)
        bold = doc.add_paragraph()
        run = bold.add_run("Week 1:")
        run.bold = True
        doc.add_paragraph("Install QDK and run Q# on the local simulator. Qubit superposition.")
        doc.add_heading("Phase 2 (Weeks 3–4): Simulator-Only Project", level=2)
        doc.add_paragraph("Week 3 – Design: Plan Grover search; no real hardware.")
        doc.add_paragraph("Week 4 – Implement: Code oracle and diffusion.")
        doc.add_heading("Phase 3 (Weeks 5–6): Test and Document", level=2)
        doc.add_paragraph("Week 5 – Testing: Validate probabilities on simulator.")
        doc.add_paragraph("Week 6 – Documentation: Write README and final report.")
    doc.save(path)


@pytest.fixture
def source_dir(tmp_path: Path) -> Path:
    directory = tmp_path / "sources"
    directory.mkdir()
    for route, _, filename in DOCUMENT_DEFINITIONS:
        make_doc(directory / filename, route)
    return directory


@pytest.fixture
def test_settings(tmp_path: Path, source_dir: Path) -> Settings:
    base = Settings.load()
    data_dir = tmp_path / "data"
    return replace(
        base,
        project_root=tmp_path,
        data_dir=data_dir,
        db_path=data_dir / "test.db",
        index_dir=data_dir / "index",
        log_path=tmp_path / "test.log",
        source_paths=tuple(source_dir / filename for _, _, filename in DOCUMENT_DEFINITIONS),
        embedding_backend="hashing",
        embedding_model="hashing-test-v1",
        chunk_max_chars=600,
        chunk_overlap_paragraphs=1,
        max_context_chars=4000,
    )


@pytest.fixture
def database(test_settings: Settings) -> Database:
    db = Database(test_settings.db_path)
    db.migrate()
    return db


@pytest.fixture
def profile_data() -> dict:
    today = date.today()
    return {
        "name": "Test Öğrencisi",
        "goal": "Pandas öğrenmek",
        "preferred_route": "data-analysis",
        "weekly_days": 5,
        "daily_minutes": 60,
        "preferred_days": ["Pazartesi", "Salı", "Çarşamba", "Perşembe", "Cuma"],
        "start_date": today.isoformat(),
        "target_end_date": (today + timedelta(days=60)).isoformat(),
        "theory_practice_preference": "Dengeli",
        "completed_topics": [],
        "difficult_topics": [],
        "review_preference": "Haftalık",
        "declared_level": "Başlangıç",
    }


@pytest.fixture
def profile_id(database: Database, profile_data: dict) -> int:
    return database.create_profile(profile_data, {"python": "Başlangıç", "pandas": "Yok"})
